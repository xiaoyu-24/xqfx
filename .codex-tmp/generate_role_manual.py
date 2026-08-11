from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"D:\projects\xqfx\docs\需求收集平台三角色使用手册.docx")

NAVY = "17324D"
BLUE = "1677FF"
CYAN = "13A8A8"
GREEN = "389E0D"
ORANGE = "D46B08"
RED = "CF1322"
PURPLE = "722ED1"
INK = "1F2937"
MUTED = "667085"
LINE = "D9E2EC"
PALE_BLUE = "EAF3FF"
PALE_GREEN = "EDF8E8"
PALE_ORANGE = "FFF4E5"
PALE_RED = "FFF1F0"
PALE_GRAY = "F5F7FA"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges:
            continue
        tag = "w:" + edge
        edge_el = tc_borders.find(qn(tag))
        if edge_el is None:
            edge_el = OxmlElement(tag)
            tc_borders.append(edge_el)
        for key, value in edges[edge].items():
            edge_el.set(qn("w:" + key), str(value))


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def set_run_font(run, size=None, bold=None, color=None, east_asia="Microsoft YaHei") -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text: str, *, bold=False, color=INK, size=None) -> None:
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def clear_cell(cell) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)


def add_cell_text(cell, text: str, *, bold=False, color=INK, size=9.5, align=None) -> None:
    clear_cell(cell)
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    add_text(p, text, bold=bold, color=color, size=size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    border = {"val": "single", "sz": "4", "color": LINE}
    set_cell_border(cell, top=border, left=border, bottom=border, right=border)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 14),
        ("Subtitle", 13, MUTED, 0, 10),
        ("Heading 1", 16, NAVY, 10, 8),
        ("Heading 2", 13, BLUE, 9, 5),
        ("Heading 3", 12, INK, 7, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True if "Heading" in style_name or style_name == "Title" else False
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(4)

    if "Callout" not in doc.styles:
        style = doc.styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        style.font.size = Pt(10)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Inches(0.12)
        style.paragraph_format.right_indent = Inches(0.08)
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(4)


def configure_headers_footers(doc: Document) -> None:
    for section in doc.sections:
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        add_text(p, "需求收集平台  |  三角色使用手册", color=MUTED, size=8.5)
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_text(p, "内部使用  ·  ", color=MUTED, size=8.5)
        add_field(p, "PAGE")


def add_paragraph(doc, text="", style=None, *, bold=False, color=INK, size=None, align=None, keep=False):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if keep:
        p.paragraph_format.keep_together = True
    if text:
        add_text(p, text, bold=bold, color=color, size=size)
    return p


def add_bullets(doc, items: list[str], numbered=False) -> None:
    style = "List Number" if numbered else "List Bullet"
    num_id = None
    if numbered:
        numbering = doc.part.numbering_part.element
        style_num_id = doc.styles["List Number"].element.pPr.numPr.numId.val
        base_num = next(node for node in numbering.findall(qn("w:num")) if int(node.get(qn("w:numId"))) == style_num_id)
        abstract_num_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
        existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
        num_id = max(existing_ids, default=0) + 1
        new_num = OxmlElement("w:num")
        new_num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), abstract_num_id)
        new_num.append(abstract_ref)
        level_override = OxmlElement("w:lvlOverride")
        level_override.set(qn("w:ilvl"), "0")
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), "1")
        level_override.append(start_override)
        new_num.append(level_override)
        numbering.append(new_num)
    for item in items:
        p = add_paragraph(doc, item, style=style)
        if num_id is not None:
            num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = num_id


def add_callout(doc, title: str, body: str, kind="info") -> None:
    palette = {
        "info": (BLUE, PALE_BLUE),
        "success": (GREEN, PALE_GREEN),
        "warning": (ORANGE, PALE_ORANGE),
        "danger": (RED, PALE_RED),
    }
    accent, fill = palette[kind]
    table = doc.add_table(rows=1, cols=2)
    set_table_width(table, [180, 9180])
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_border(table.cell(0, 0))
    set_cell_border(table.cell(0, 1))
    table.cell(0, 0).width = Inches(0.125)
    p = table.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    add_text(p, title, bold=True, color=accent, size=10)
    p2 = table.cell(0, 1).add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    add_text(p2, body, color=INK, size=9.5)
    set_cell_margins(table.cell(0, 1), top=100, start=140, bottom=100, end=140)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int], header_fill=NAVY, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_width(table, widths)
    header = table.rows[0]
    repeat_table_header(header)
    prevent_row_split(header)
    for i, value in enumerate(headers):
        add_cell_text(header.cells[i], value, bold=True, color=WHITE, size=9.2)
        set_cell_shading(header.cells[i], header_fill)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, value in enumerate(values):
            add_cell_text(cells[i], value, size=font_size)
            if row_idx % 2 == 1:
                set_cell_shading(cells[i], PALE_GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_flow(doc, steps: list[tuple[str, str]], color=BLUE) -> None:
    widths = [1120] * (len(steps) - 1) + [9360 - 1120 * (len(steps) - 1)]
    table = doc.add_table(rows=2, cols=len(steps))
    set_table_width(table, widths)
    for i, (number, label) in enumerate(steps):
        cell = table.cell(0, i)
        set_cell_shading(cell, color)
        add_cell_text(cell, number, bold=True, color=WHITE, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, color)
        body = table.cell(1, i)
        add_cell_text(body, label, bold=True, color=INK, size=8.6, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(body, "F8FAFC")
    for row in table.rows:
        prevent_row_split(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_role_cover(doc, index: str, title: str, subtitle: str, color: str, daily: str) -> None:
    if doc.paragraphs and not doc.paragraphs[-1].text and not doc.paragraphs[-1]._p.xpath(".//w:br"):
        trailing = doc.paragraphs[-1]._element
        trailing.getparent().remove(trailing)
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(96)
    p.paragraph_format.space_after = Pt(8)
    add_text(p, index, bold=True, color=color, size=12)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    add_text(p, title, bold=True, color=NAVY, size=28)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(30)
    add_text(p, subtitle, color=MUTED, size=13)
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [9360])
    set_cell_shading(table.cell(0, 0), color)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    add_text(p, "每日工作主线", bold=True, color=WHITE, size=10)
    p2 = table.cell(0, 0).add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    add_text(p2, daily, color=WHITE, size=12)
    set_cell_margins(table.cell(0, 0), top=220, start=260, bottom=220, end=260)


def add_page_break(doc) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build_document() -> Document:
    doc = Document()
    style_document(doc)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(75)
    p.paragraph_format.space_after = Pt(10)
    add_text(p, "需求收集平台", bold=True, color=BLUE, size=14)
    p = doc.add_paragraph(style="Title")
    add_text(p, "三角色使用手册", bold=True, color=NAVY, size=30)
    p = doc.add_paragraph(style="Subtitle")
    add_text(p, "普通用户 · 需求处理员 · 管理员", color=MUTED, size=13)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    role_table = doc.add_table(rows=1, cols=3)
    set_table_width(role_table, [3120, 3120, 3120])
    for i, (code, role, desc, color) in enumerate([
        ("USER", "普通用户", "提交与跟进自己的需求", BLUE),
        ("HANDLER", "需求处理员", "评估、推进与版本协同", CYAN),
        ("ADMIN", "管理员", "治理账号、字典与 AI", PURPLE),
    ]):
        cell = role_table.cell(0, i)
        set_cell_shading(cell, color)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, code, bold=True, color=WHITE, size=9)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p2, role, bold=True, color=WHITE, size=14)
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(0)
        add_text(p3, desc, color=WHITE, size=8.5)
        set_cell_margins(cell, top=210, start=120, bottom=210, end=120)
        prevent_row_split(role_table.rows[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(6)
    add_text(p, "文档定位", bold=True, color=NAVY, size=11)
    add_paragraph(doc, "本手册按岗位拆分每日工作和标准操作，可直接用于新员工上手、岗位交接与日常自查。内容依据当前系统界面、路由权限和服务端规则整理。", color=MUTED, size=10)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    add_text(p, f"版本：V1.0  |  编制日期：{date.today().isoformat()}  |  适用环境：当前部署版本", color=MUTED, size=9)

    # Quick start
    add_page_break(doc)
    doc.add_heading("快速开始", level=1)
    add_paragraph(doc, "先找到自己的角色，再按对应章节执行。三个角色都需要完成登录、首次改密、消息查看与退出登录。")

    doc.add_heading("1. 角色与入口", level=2)
    add_table(doc, ["角色", "进入系统后优先打开", "核心产出", "本手册章节"], [
        ["普通用户 USER", "需求列表 / 填写需求", "信息完整、可验收的需求", "第一篇"],
        ["需求处理员 HANDLER", "待办工作台", "持续更新的处理进度和状态", "第二篇"],
        ["管理员 ADMIN", "待办工作台；按需进入管理菜单", "可用的账号、字典、系统和 AI 配置", "第三篇"],
    ], [1800, 2300, 3200, 2060])

    doc.add_heading("2. 全员首次登录", level=2)
    add_flow(doc, [("01", "管理员分配账号"), ("02", "输入账号和密码"), ("03", "首次登录修改密码"), ("04", "使用新密码重新登录"), ("05", "进入角色首页")])
    add_bullets(doc, [
        "登录账号由管理员分配；忘记密码时联系管理员重置。",
        "首次登录或密码被重置后，系统强制修改密码；新密码至少 8 位。",
        "改密成功后全部登录会话失效，必须用新密码重新登录。",
        "右上角显示当前姓名和角色；离开公共电脑前务必点击“退出”。",
    ])
    add_callout(doc, "账号安全", "管理员创建或重置账号后，页面会显示初始密码为 888888。初始密码仅用于首次登录，应立即修改，不要在群聊或文档中长期保存。", "warning")

    doc.add_heading("3. 全局导航和站内消息", level=2)
    add_table(doc, ["位置", "用途", "操作要点"], [
        ["左侧菜单", "进入当前角色可用页面", "菜单根据角色自动显示；无权限页面会被拦截"],
        ["顶部搜索", "按关键词搜索需求", "输入标题或内容关键词后进入需求列表"],
        ["通知铃铛", "查看未读数量和站内消息", "支持全部/仅未读、单条已读、全部已读"],
        ["通知记录", "跳转到关联需求", "点击后先标为已读，再打开需求详情"],
        ["退出", "结束当前会话", "填写页有未保存内容时会先确认是否离开"],
    ], [1600, 2500, 5260])
    add_callout(doc, "通知接收规则", "普通用户需求的“状态更新”和“进度更新”始终发给实际登录并提交该需求的账号，不按表单里填写的姓名寻找收件人。新需求、超期和长期无进展提醒发给所属系统负责人及协助人。", "info")

    doc.add_heading("4. 需求生命周期", level=2)
    add_flow(doc, [("01", "草稿（可选）"), ("02", "待评估"), ("03", "已确认"), ("04", "开发中 / 暂停"), ("05", "已完成 / 已拒绝 / 已关闭")], color=CYAN)
    add_paragraph(doc, "状态由需求处理员或管理员在“更新进度”时修改；编辑需求内容时不能直接改状态。状态并非固定单向按钮，处理人员应根据真实业务结果选择，并在进展内容中写清原因。", color=MUTED, size=9.5)

    doc.add_heading("5. 快速导航", level=2)
    add_table(doc, ["我要做什么", "去哪里", "看哪一节"], [
        ["提交、暂存或继续完善需求", "填写需求 / 需求列表", "普通用户 2-4"],
        ["处理今日待办", "待办工作台", "需求处理员 2"],
        ["把需求纳入版本", "系统与版本 > 管理需求", "需求处理员 5"],
        ["管理账号、字典或 AI", "对应的管理员菜单", "管理员 2-4"],
    ], [3400, 3000, 2960])

    # USER
    add_role_cover(doc, "第一篇 · USER", "普通用户操作手册", "目标：准确提交需求，并持续跟进属于自己账号的处理结果。", BLUE,
                   "查看消息 → 跟进已有需求 → 补充/编辑允许修改的需求 → 提交新的需求")
    doc.add_heading("1. 岗位职责和权限边界", level=1)
    add_table(doc, ["可以做", "不可以做"], [[
        "填写、暂存和提交需求；查看自己账号提交的需求；在允许状态下编辑或删除；查看详情、进展、版本记录和附件；查看系统与版本只读信息；接收自己的状态/进度通知。",
        "查看他人需求；进入待办工作台或需求概览；更新处理进度和状态；管理版本需求；管理人员、字典或 AI 配置。",
    ]], [4680, 4680])
    add_callout(doc, "数据归属", "系统以“实际登录提交需求的账号”判断归属。表单中的“姓名”是业务展示字段，即使填写成其他姓名，也不会改变需求归属和通知接收账号。", "warning")

    doc.add_heading("2. 每日工作清单", level=1)
    add_table(doc, ["时间", "动作", "完成标准"], [
        ["上班后", "点击右上角通知铃铛，先看未读状态/进度消息", "重要变化已打开并理解；已处理消息已读"],
        ["需要跟进时", "打开需求详情查看最新进展、状态和版本记录", "确认是否需要补充业务信息或验收反馈"],
        ["产生新需求时", "先整理背景、问题、期望结果、验收标准和周期", "提交内容能被处理人员独立理解"],
        ["提交后", "回到需求列表确认状态和附件", "正式需求显示“待评估”，附件可下载"],
        ["下班前", "检查草稿和待评估需求", "草稿已继续完善；无重复或误提需求"],
    ], [1300, 3800, 4260])

    doc.add_heading("3. 新建需求：标准流程", level=1)
    add_flow(doc, [("01", "打开“填写需求”"), ("02", "可选：AI 分析并填充"), ("03", "核对必填和周期"), ("04", "选择/拖入附件"), ("05", "暂存或保存需求")], color=BLUE)
    doc.add_heading("3.1 手工填写", level=2)
    add_table(doc, ["字段", "要求", "建议写法"], [
        ["姓名 *", "默认读取当前账号姓名，可按业务需要修正", "填写实际需求联系人"],
        ["部门 *", "从管理员维护的启用部门中选择", "选择需求归属部门"],
        ["需求标题 *", "简要概括；列表和通知都使用该标题", "动词 + 对象 + 期望结果"],
        ["类型 *", "从启用的需求类型中选择", "选择最接近的业务分类"],
        ["紧急程度 *", "高 / 中 / 低", "按业务影响和时限判断，不把所有需求都设为高"],
        ["时间周期", "选填；开始和结束必须同时填写，结束不得早于开始", "写真实期望窗口"],
        ["所属系统 *", "只能选择启用中的系统", "选错系统会影响负责人收到通知"],
        ["需求内容 *", "写清背景、问题、期望结果和验收标准", "用短段落或编号描述"],
        ["附件", "图片/PDF/Word/Excel；单个最大 100MB", "只上传与需求相关的最新版材料"],
    ], [1600, 3450, 4310], font_size=8.8)
    add_callout(doc, "正式提交前自检", "标题能否一句话说明问题？所属系统是否正确？需求内容是否包含验收标准？周期是否成对填写？附件是否为最终版本？", "success")

    doc.add_heading("3.2 使用 AI 智能导入", level=2)
    add_bullets(doc, [
        "展开“AI 智能导入”，粘贴一段需求描述，点击“分析并填充”。",
        "AI 只填充当前为空的字段，不覆盖已经填写的内容，也不会自动提交。",
        "分析结果必须人工核对，重点检查部门、类型、所属系统、日期和验收标准。",
        "若提示未识别有效字段，可改用结构化文本重新分析，或直接手工填写。",
        "若提示 AI 分析失败，保留当前表单内容，改为手工填写并联系管理员检查 AI 配置。",
    ], numbered=True)
    add_callout(doc, "推荐输入结构", "背景：……；当前问题：……；期望结果：……；涉及系统：……；计划周期：……；验收标准：……。", "info")

    doc.add_heading("3.3 暂存与正式保存", level=2)
    add_table(doc, ["动作", "适用场景", "结果"], [
        ["暂存", "信息尚未收齐", "允许缺少正式必填项；生成草稿并留在填写流程中"],
        ["保存需求", "信息已完整", "执行完整校验；成功后进入需求列表，状态为“待评估”"],
        ["离开未保存页面", "切换菜单或退出", "系统提示内容尚未保存，可选择继续编辑或离开"],
    ], [1700, 3200, 4460])

    doc.add_heading("4. 附件操作", level=1)
    add_bullets(doc, [
        "点击上传区选择多个文件，或把文件拖入上传区；重复文件会自动去重。",
        "支持 JPG/JPEG、PNG、GIF、WEBP、PDF、DOC/DOCX、XLS/XLSX；单个文件最大 100MB。",
        "图片和 PDF 可直接预览；Word/Excel 需要后台生成预览，可能显示等待中或转换中。",
        "预览失败可点击“重新生成”；无论预览是否成功，都可下载原文件。",
        "在线预览可能与原文件略有差异，正式内容以下载的原文件为准。",
    ], numbered=True)

    doc.add_heading("5. 查询、跟进和修改", level=1)
    add_heading = doc.add_heading
    add_heading("5.1 需求列表", level=2)
    add_paragraph(doc, "普通用户的列表只显示当前登录账号提交的需求。默认打开“未完结”，需要找草稿或终态需求时切换到“全部”。")
    add_table(doc, ["操作", "使用方法"], [
        ["关键词搜索", "搜索标题或需求内容；顶部全局搜索也会进入该列表"],
        ["组合筛选", "按状态、所属系统、目标版本、部门、填写人、类型、保存类型组合筛选"],
        ["清除筛选", "点击筛选标签单独清除，或点击“清空全部”"],
        ["查看详情", "查看需求内容、状态、进展时间线、版本变更和附件"],
        ["编辑", "仅自己的草稿、待评估或已确认需求可编辑"],
        ["删除", "仅自己的草稿或待评估需求可删除；删除前会二次确认"],
    ], [2000, 7360])

    doc.add_heading("5.2 消息与进展", level=2)
    add_bullets(doc, [
        "收到“状态更新”或“进度更新”后，点击通知直接进入需求详情。",
        "详情页按时间显示处理人、进展内容和新状态；版本绑定、迁移或解除也有单独记录。",
        "普通用户只能查看进展，不能新增进展或修改状态。",
        "如果业务信息发生变化，在允许编辑的状态下修改需求；若已进入开发中或终态，联系处理人员通过进展记录沟通。",
    ])

    doc.add_heading("6. 常见问题", level=1)
    add_table(doc, ["现象", "处理方法"], [
        ["看不到某条需求", "确认是否由当前登录账号实际提交；普通用户不能查看他人需求"],
        ["不能编辑", "确认是否为自己的需求，且状态是否为草稿、待评估或已确认"],
        ["不能删除", "只有自己的草稿或待评估需求可删除"],
        ["所属系统没有选项", "系统可能已停用；联系处理员或管理员确认"],
        ["附件预览一直等待", "稍后再看；仍失败时重新生成或下载原文件"],
        ["提示记录已被修改", "重新打开详情获取最新数据，再次操作，避免覆盖他人修改"],
        ["没有收到进度通知", "确认需求由当前账号提交，并检查“全部消息”而非仅未读；仍异常联系管理员"],
    ], [2700, 6660])

    # HANDLER
    add_role_cover(doc, "第二篇 · HANDLER", "需求处理员操作手册", "目标：让每条负责或协助的需求都有明确状态、最新进展和可追踪版本。", CYAN,
                   "看待办与提醒 → 评估高优先级需求 → 更新进度/状态 → 维护系统与版本 → 日终清零异常")
    doc.add_heading("1. 岗位职责和权限边界", level=1)
    add_table(doc, ["拥有能力", "不包含"], [[
        "普通用户全部能力；查看全部需求；使用待办工作台和需求概览；编辑/删除需求；新增处理进度并更新状态；维护系统、负责人、协助人和版本；集中绑定、迁移或解除版本需求。",
        "人员管理；部门/需求类型字典管理；AI 配置管理。",
    ]], [6200, 3160])
    add_callout(doc, "责任分配现状", "当前前端没有独立“手工指派处理人”入口。工作台和提醒主要依据系统负责人、协助人组织；因此所属系统及其负责人配置是否准确，直接决定谁看到待办和提醒。", "warning")

    doc.add_heading("2. 每日工作清单", level=1)
    add_table(doc, ["时间", "动作", "完成标准"], [
        ["上班后", "打开通知铃铛，处理新需求、超期、无进展和待处理消息", "紧急消息已打开；明确今日优先级"],
        ["上午", "检查“我负责的系统需求”和“我协助处理的”", "新需求已评估；高紧急度先处理"],
        ["处理中", "每次有实质变化即更新进度；必要时同步状态", "进展可复述已做、结果、下一步和时间点"],
        ["版本规划时", "检查目标版本，纳入未分配需求或迁移其他版本需求", "版本范围与交付计划一致"],
        ["下班前", "检查未完结、超期、长期无进展项", "当天有动作的需求已留痕；无“口头更新但系统无记录”"],
    ], [1300, 4100, 3960])

    doc.add_heading("3. 发现和判断待处理需求", level=1)
    doc.add_heading("3.1 待办工作台", level=2)
    add_bullets(doc, [
        "“我负责的系统需求”展示当前账号作为系统负责人的全部未完结需求。",
        "“我协助处理的”展示当前账号作为系统协助人的未完结需求；无协助项时该区域不显示。",
        "每条需求显示标题、系统、提出人、更新时间、紧急程度和状态；点击即可进入详情。",
        "先看高紧急度和长时间未更新项，再处理普通新增需求。",
    ])
    doc.add_heading("3.2 需求概览", level=2)
    add_table(doc, ["区域", "用途"], [
        ["需求总数 / 草稿", "了解全平台规模与未正式提交数量"],
        ["未完结", "查看待评估与处理中总量"],
        ["已完成 / 完结率", "观察交付结果"],
        ["高紧急度待处理", "确定优先跟进对象"],
        ["按系统分布", "点击图例跳转到该系统需求列表"],
        ["按状态 / 紧急度分布", "识别积压阶段和风险结构"],
    ], [2800, 6560])
    doc.add_heading("3.3 全部需求列表", level=2)
    add_paragraph(doc, "默认查看未完结，可切换全部。处理员能查看全平台需求，并按关键词、系统、版本、部门、填写人、类型、状态、保存类型组合筛选。终态需求要先切换到“全部”。")

    doc.add_heading("4. 处理一条需求", level=1)
    add_flow(doc, [("01", "打开详情并核对"), ("02", "确认业务与附件"), ("03", "补充进展内容"), ("04", "可选：同步状态"), ("05", "提交并复核时间线")], color=CYAN)
    add_bullets(doc, [
        "核对标题、所属系统、负责人、目标版本、提出人、周期、需求内容和附件。",
        "需要修改需求本身时点击“编辑需求”；内容编辑页的状态字段只读。",
        "点击“更新进度”，填写当前进展、处理结果或补充说明。进展内容必填。",
        "“更新后状态”可选；只记录过程时留空，确有阶段变化时选择新状态。",
        "提交后详情页新增时间线记录；若状态变化，实际提交该需求的普通用户收到状态/进度通知。",
    ], numbered=True)
    add_callout(doc, "高质量进展模板", "已完成：……；当前结果/风险：……；下一步：……；预计时间：……；需要提出人配合：……。", "success")

    doc.add_heading("4.1 状态使用建议", level=2)
    add_table(doc, ["状态", "何时使用", "进展里至少写清"], [
        ["待评估", "刚提交、尚未完成判断", "待确认的问题和负责人"],
        ["已确认", "范围、目标和受理结论已明确", "确认内容、初步计划"],
        ["开发中", "已进入实施", "当前阶段、下一里程碑"],
        ["暂停", "暂时不能继续", "暂停原因、恢复条件和复查日期"],
        ["已完成", "已交付或已达到验收条件", "交付结果和验证方式"],
        ["已拒绝", "明确不受理", "拒绝依据和替代建议"],
        ["已关闭", "无需继续推进", "关闭原因和确认情况"],
    ], [1450, 3350, 4560], font_size=8.8)

    doc.add_heading("5. 系统和版本协同", level=1)
    doc.add_heading("5.1 系统维护", level=2)
    add_bullets(doc, [
        "进入“系统与版本”，可按系统名称、负责人、协助人和状态筛选。",
        "新增/编辑系统时，系统名称和负责人必填；协助人可多选。负责人和协助人必须是启用的处理员或管理员。",
        "负责人不能同时作为同一系统协助人，协助账号不能重复。",
        "停用系统后不能再用于新建需求，但历史需求仍保留。",
        "删除有关联需求的系统时，系统会要求先迁移到其他启用系统或“暂无系统”。",
    ], numbered=True)
    doc.add_heading("5.2 版本维护", level=2)
    add_bullets(doc, [
        "先选择左侧系统，再新增或编辑版本；版本名称必填，说明最多 2000 字。",
        "停用版本后不能纳入新需求，已有需求仍可查看和解除。",
        "删除有关联需求的版本会失败；先进入“管理需求”解除或迁移。",
    ])

    doc.add_heading("5.3 集中管理版本需求", level=2)
    add_table(doc, ["区域", "操作", "规则"], [
        ["当前版本需求", "按标题、状态、紧急度筛选；批量解除绑定", "已完成/已拒绝/已关闭等终态只读，不能迁移或解除"],
        ["可纳入需求", "从未绑定版本或其他版本中选择，点击“纳入当前版本”", "当前版本必须启用；系统自动区分绑定和迁移数量"],
        ["确认弹窗", "核对本次绑定、迁移或解除的数量", "操作立即生效，并写入需求的版本变更记录"],
    ], [1900, 3700, 3760])

    doc.add_heading("6. 提醒与异常处置", level=1)
    add_table(doc, ["通知/现象", "含义", "处理动作"], [
        ["新需求", "所属系统有正式需求提交", "打开详情完成初步评估并留进展"],
        ["待处理", "后端已将需求指派到当前账号", "打开详情并及时处理；当前前端无独立指派入口"],
        ["已超期", "需求周期结束日已过且未终结", "确认是否调整计划或推进终态，并记录原因"],
        ["无进展", "达到系统配置的连续无进展天数", "补充真实进度；不要为了消除提醒写空洞信息"],
        ["并发修改提示", "他人在你打开页面后更新了记录", "重新打开/刷新获取最新版本，再操作"],
        ["版本无法删除", "仍有关联需求", "先管理版本需求，解除或迁移后再删"],
    ], [1900, 3300, 4160])

    # ADMIN
    add_role_cover(doc, "第三篇 · ADMIN", "管理员操作手册", "目标：保证账号、基础字典、系统责任和 AI 服务长期可用且可追踪。", PURPLE,
                   "先处理业务待办 → 检查账号与基础数据 → 维护系统责任 → 核对 AI 当前配置 → 处理用户反馈")
    doc.add_heading("1. 管理员职责和每日节奏", level=1)
    add_paragraph(doc, "管理员拥有需求处理员的全部业务能力，并额外负责人员、字典和 AI 配置。日常应先完成需求处理员工作，再处理治理任务。")
    add_table(doc, ["频率", "工作", "完成标准"], [
        ["每天", "查看待办/通知；处理账号启停、忘记密码、角色变更请求", "无积压的登录阻塞；高风险变更已核对"],
        ["每周", "检查长期无进展、超期需求；核对系统负责人和协助人", "离职/转岗账号不再承担系统责任"],
        ["按需", "新增人员、部门、需求类型、系统、版本或 AI 配置", "名称清晰、无重复、责任人有效"],
        ["变更后", "使用实际账号验证菜单和关键流程", "权限、通知和表单选项符合预期"],
    ], [1400, 4200, 3760])

    doc.add_heading("2. 人员管理", level=1)
    doc.add_heading("2.1 新增和编辑账号", level=2)
    add_bullets(doc, [
        "进入“人员管理”，可按账号/姓名、角色、状态筛选。",
        "点击“新增人员”，填写账号、姓名、部门和角色；账号创建后不可修改。",
        "新建成功后页面显示初始密码 888888；通过安全渠道单独告知本人。",
        "编辑账号可调整姓名、部门和角色；角色变更从该账号下一次请求起生效。",
    ], numbered=True)
    add_table(doc, ["角色", "授予前确认"], [
        ["普通用户", "只需提交和跟进自己的需求；不能作为系统负责人或协助人"],
        ["需求处理员", "需要查看全部需求、更新进度、维护系统和版本"],
        ["管理员", "确实需要人员、字典和 AI 配置权限；严格控制数量"],
    ], [2000, 7360])
    add_callout(doc, "降级检查", "把处理员或管理员降级为普通用户前，先确认其没有仍绑定的系统负责人或协助人职责，否则保存可能被服务端拦截。", "warning")

    doc.add_heading("2.2 重置密码和启停账号", level=2)
    add_bullets(doc, [
        "重置前确认用户身份；点击“重置密码”并二次确认。",
        "重置后该用户原有全部登录会话失效，初始密码恢复为 888888，首次登录必须改密。",
        "停用账号用于离职、长期离岗或安全事件；停用后无法登录，也不能被新选为负责人或协助人。",
        "重新启用前核对角色、部门和系统责任是否仍正确。",
    ], numbered=True)

    doc.add_heading("3. 字典管理", level=1)
    add_flow(doc, [("01", "选择部门/需求类型"), ("02", "新增或改名"), ("03", "确认表单选项"), ("04", "不再使用则停用"), ("05", "必要时重新启用")], color=PURPLE)
    add_bullets(doc, [
        "字典包含“部门”和“需求类型”两个分类，名称最长 50 字。",
        "新增或改名后，填写需求、编辑需求、人员部门和 AI 回填选项会使用最新启用数据。",
        "停用项不会出现在新建表单和 AI 回填中，但历史需求仍保留原值。",
        "不要通过改名把一个业务含义变成完全不同的含义；这会让历史统计难以解释。新含义应新增条目，旧条目停用。",
    ])

    doc.add_heading("4. AI 配置", level=1)
    add_table(doc, ["字段/动作", "说明", "注意事项"], [
        ["配置名称", "识别用途，如“生产主模型”", "名称要能区分环境和供应商"],
        ["AI Base URL", "模型服务基础地址", "确认协议、路径和网络可达"],
        ["模型名称", "智能分析调用的模型标识", "必须与服务端实际可用模型一致"],
        ["API Key", "模型服务访问凭据", "页面只显示掩码；编辑时留空表示保留旧 Key，输入新值才替换"],
        ["启用/停用", "控制配置是否可用", "停用前确认已有可用替代配置"],
        ["切换使用", "把该配置设为“当前使用”", "同一时间只有一个当前使用配置"],
    ], [1900, 3300, 4160])
    add_bullets(doc, [
        "点击“新增 AI 配置”，完整填写地址、模型和 Key，保存后根据需要点击“切换使用”。",
        "编辑配置时，API Key 输入框为空不会清除旧值；只有输入新值才会替换。",
        "“当前使用”配置不能再次切换；若需停用，先准备并切换到可用配置。",
        "保存成功不等于智能导入一定可用。使用普通表单执行一次“分析并填充”作为业务验证。",
        "出现“系统处理失败，请联系管理员并提供追踪编号”时，记录追踪编号、发生时间和所用配置，检查后端日志与 AI 地址/模型/Key。",
    ], numbered=True)
    add_callout(doc, "凭据管理", "不要把 API Key 写入聊天记录、截图、手册或工单正文。只在 AI 配置表单内录入；需要轮换时直接输入新 Key 覆盖旧值。", "danger")

    doc.add_heading("5. 系统与版本治理", level=1)
    add_bullets(doc, [
        "管理员应定期核对每个系统的负责人、协助人和状态，尤其是人员离职或角色变更后。",
        "删除系统前迁移关联需求；迁移到其他系统时目标必须启用，也可迁移到“暂无系统”。",
        "停用版本不再接收新纳入需求；终态需求在版本管理页只读。",
        "系统名称和同一系统内版本名称保持唯一、稳定、可识别。",
    ])

    doc.add_heading("6. 管理变更后的验证清单", level=1)
    add_table(doc, ["变更", "最小验证"], [
        ["新增/改角色", "用目标账号重新登录，确认菜单、数据范围和接口权限"],
        ["重置密码", "确认旧会话失效，新初始密码要求改密，新密码可登录"],
        ["停用账号", "确认不能登录，且不再出现在负责人/协助人可选项"],
        ["字典变更", "打开填写需求，确认启用项出现、停用项消失、历史需求仍正常"],
        ["系统责任变更", "用新负责人/协助人查看工作台和新需求通知"],
        ["AI 配置变更", "执行一次智能导入，核对只填空字段且不自动提交"],
    ], [2300, 7060])

    # Appendices
    add_page_break(doc)
    doc.add_heading("附录 A · 权限矩阵", level=1)
    add_table(doc, ["功能", "普通用户", "需求处理员", "管理员"], [
        ["填写/暂存需求", "可", "可", "可"],
        ["查看需求", "仅自己账号提交", "全部", "全部"],
        ["编辑需求", "自己的草稿/待评估/已确认", "全部", "全部"],
        ["删除需求", "自己的草稿/待评估", "可", "可"],
        ["查看进展和附件", "仅自己的需求", "全部", "全部"],
        ["新增进展/更新状态", "不可", "可", "可"],
        ["待办工作台/概览", "不可", "可", "可"],
        ["系统与版本", "只读", "维护", "维护"],
        ["版本需求集中管理", "不可", "可", "可"],
        ["站内消息", "可", "可", "可"],
        ["人员/字典/AI 配置", "不可", "不可", "可"],
    ], [3300, 2020, 2020, 2020], font_size=8.5)

    doc.add_heading("附录 B · 通知路由表", level=1)
    add_table(doc, ["通知类型", "触发条件", "接收人", "点击后"], [
        ["新需求", "正式需求提交", "所属系统负责人和协助人；不重复通知提交人", "需求详情"],
        ["待处理", "需求被后端指派", "被指派处理人；不重复通知操作人", "需求详情"],
        ["状态更新", "更新进度时同步改变状态", "实际登录提交该需求的账号；不重复通知操作人", "需求详情"],
        ["进度更新", "新增进展但状态未变", "实际登录提交该需求的账号；不重复通知操作人", "需求详情"],
        ["已超期", "未终结需求的周期结束日已过", "所属系统负责人和协助人", "需求详情"],
        ["无进展", "未终结需求达到连续无进展阈值", "所属系统负责人和协助人", "需求详情"],
    ], [1550, 2350, 3330, 2130], font_size=8.4)

    doc.add_heading("附录 C · 错误与求助信息", level=1)
    add_table(doc, ["提示/问题", "先做什么", "仍未解决时提供"], [
        ["系统处理失败，请联系管理员并提供追踪编号", "记录编号，不反复提交；确认发生在哪个页面和动作", "追踪编号、时间、账号、页面、动作；不要提供密码或 Key"],
        ["保存失败/必填项错误", "查看红色字段提示，补齐后再保存", "字段名、提示文案、需求编号"],
        ["记录已被其他人修改", "重新打开页面获取最新数据", "需求编号、发生时间、操作动作"],
        ["无权访问/无权操作", "确认右上角角色和数据归属", "账号、角色、目标需求编号"],
        ["AI 分析失败", "保留表单，改用手工填写；确认当前 AI 配置", "追踪编号、配置名称、时间；禁止提供完整 Key"],
        ["附件无法预览", "下载原文件；失败状态可点重新生成", "需求编号、附件名、预览状态"],
    ], [2900, 3320, 3140], font_size=8.5)
    add_callout(doc, "禁止提交的敏感信息", "任何求助材料都不要包含登录密码、数据库密码、完整 API Key、会话令牌或浏览器 Cookie。", "danger")

    doc.add_heading("附录 D · 新用户 15 分钟上手", level=1)
    add_table(doc, ["分钟", "练习", "验收"], [
        ["0-3", "登录并完成首次改密；识别右上角姓名、角色、通知和退出", "能安全进入系统"],
        ["3-7", "打开需求列表，切换未完结/全部并尝试筛选", "能找到目标需求"],
        ["7-12", "新建一个草稿；体验 AI 只填空字段和附件选择", "草稿可在列表中找到"],
        ["12-15", "打开草稿详情，编辑并了解状态/进展区域", "知道角色允许的下一步"],
    ], [1100, 5000, 3260])

    configure_headers_footers(doc)
    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
